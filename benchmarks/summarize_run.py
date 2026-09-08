#!/usr/bin/env python3
"""Summarise a C-MAGE run with no ground truth: what came out, per PDF and page.

For corpora without known answers (real patents) the honest report is counts and
the confidence distribution, not an accuracy figure. This reads:

  logs/stage1.log   "PDF <stem> is loaded." then "Page <i> saved. Number of objects: <n>"
                    per page, in order -- the only record of which page a figure
                    came from, since figure names carry a running index only
  02_DIS_Segments/  segment images named Image_DIS_VH_File_<stem>_image_<n>_molecule_<k>.png
  03_CXMS_Results/  the two spreadsheets (confidence, tier, SMILES)
  --pdf-dir         page counts via pdfinfo, so pages that produced nothing are visible

Outputs in --out: summary.md, summary.json, structures.csv (file name, PDF, page,
tier, confidence, SMILES, parses-with-RDKit, heavy atoms).

  summarize_run.py --run-dir results/run_X --pdf-dir corpus/ --out results/patents [--threshold 0.8431]
"""
import argparse
import csv
import json
import re
import subprocess
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

DIS_PREFIX = "Image_DIS_VH_File_"


def figure_pages(stage1_log):
    """{pdf stem: {figure index (1-based): page number (1-based)}} from the stage-1 log."""
    out = {}
    if not stage1_log.exists():
        return out
    cur = None
    counter = 0
    for line in stage1_log.read_text(errors="replace").splitlines():
        m = re.match(r"PDF (.+) is loaded\.", line)
        if m:
            cur, counter = m.group(1), 0
            out[cur] = {}
            continue
        m = re.match(r"Page (\d+) saved\. Number of objects: (\d+)", line)
        if m and cur:
            page, n = int(m.group(1)) + 1, int(m.group(2))
            for k in range(n):
                counter += 1
                out[cur][counter] = page
    return out


def pdf_pages(pdf_dir):
    out = {}
    if not pdf_dir:
        return out
    for pdf in sorted(pdf_dir.glob("*.pdf")):
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
        m = re.search(r"^Pages:\s+(\d+)", info, re.M)
        out[pdf.stem] = int(m.group(1)) if m else None
    return out


def read_sheet(path, tier):
    rows = []
    if not path.exists():
        return rows
    ws = load_workbook(path, read_only=True).active
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r is None or len(r) < 10 or r[1] is None:
            continue
        rows.append({"file": str(r[1]), "smiles": None if r[5] is None else str(r[5]),
                     "confidence": None if r[8] is None else float(r[8]), "tier": tier})
    return rows


def split_name(path):
    stem = Path(path).stem
    if stem.startswith(DIS_PREFIX):
        stem = stem[len(DIS_PREFIX):]
    seg = fig = None
    m = re.match(r"^(.*)_molecule_(\d+)$", stem)
    if m:
        stem, seg = m.group(1), int(m.group(2))
    m = re.match(r"^(.*)_image_(\d+)$", stem)
    if m:
        stem, fig = m.group(1), int(m.group(2))
    return stem, fig, seg


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--run-dir", type=Path, required=True)
    ap.add_argument("--pdf-dir", type=Path, default=None)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--threshold", type=float, default=0.8431)
    ap.add_argument("--label", default="")
    args = ap.parse_args()

    try:
        from rdkit import Chem, RDLogger
        RDLogger.DisableLog("rdApp.*")
    except ImportError:
        Chem = None

    fp = figure_pages(args.run_dir / "logs" / "stage1.log")
    pages = pdf_pages(args.pdf_dir)
    figures = sorted(p.name for p in (args.run_dir / "01_VH_Figures").glob("*.png")) if (args.run_dir / "01_VH_Figures").exists() else []
    segments = sorted(p.name for p in (args.run_dir / "02_DIS_Segments").glob("*.png")) if (args.run_dir / "02_DIS_Segments").exists() else []
    res = args.run_dir / "03_CXMS_Results"
    rows = read_sheet(res / "Completed_HighConfidence_CMAGE.xlsx", "high") + read_sheet(res / "Completed_LowConfidence_CMAGE.xlsx", "low")

    structs = []
    for r in rows:
        pdf, fig, seg = split_name(r["file"])
        page = fp.get(pdf, {}).get(fig)
        mol = Chem.MolFromSmiles(r["smiles"]) if (Chem and r["smiles"] and r["smiles"] != "<invalid>") else None
        structs.append({"file_name": Path(r["file"]).name, "pdf": pdf, "figure": fig, "page": page, "segment": seg,
                        "tier": r["tier"], "confidence": r["confidence"], "smiles": r["smiles"],
                        "parses": mol is not None, "heavy_atoms": mol.GetNumHeavyAtoms() if mol else None})

    per_pdf = {}
    all_pdfs = sorted(set(pages) | set(fp) | {s["pdf"] for s in structs})
    for pdf in all_pdfs:
        ss = [s for s in structs if s["pdf"] == pdf]
        figs = [f for f in figures if split_name(f)[0] == pdf]
        segs = [f for f in segments if split_name(f)[0] == pdf]
        pages_with_figs = sorted(set(fp.get(pdf, {}).values()))
        by_page = Counter(s["page"] for s in ss)
        per_pdf[pdf] = {
            "pages": pages.get(pdf), "pages_with_figures": len(pages_with_figs), "figures": len(figs), "segments": len(segs),
            "structures": len(ss), "high": sum(1 for s in ss if s["tier"] == "high"), "low": sum(1 for s in ss if s["tier"] == "low"),
            "invalid": sum(1 for s in ss if not s["parses"]),
            "structures_per_page": {str(k): v for k, v in sorted(by_page.items(), key=lambda kv: (kv[0] is None, kv[0]))},
            "median_confidence": (sorted(s["confidence"] for s in ss if s["confidence"] is not None)[len(ss) // 2] if ss else None),
        }

    confs = [s["confidence"] for s in structs if s["confidence"] is not None]
    bins = [0, 0.2, 0.4, 0.6, 0.7, 0.8, args.threshold, 0.9, 0.95, 1.0001]
    hist = Counter()
    for c in confs:
        for lo, hi in zip(bins, bins[1:]):
            if lo <= c < hi:
                hist[f"[{lo:.4g},{hi if hi < 1 else 1:.4g})"] += 1
                break
    summary = {"label": args.label, "run": args.run_dir.name, "threshold": args.threshold,
               "pdfs": len(all_pdfs), "pages": sum(v for v in pages.values() if v) if pages else None,
               "figures": len(figures), "segments": len(segments), "structures": len(structs),
               "high": sum(1 for s in structs if s["tier"] == "high"), "low": sum(1 for s in structs if s["tier"] == "low"),
               "invalid": sum(1 for s in structs if not s["parses"]),
               "confidence_histogram": {k: hist[k] for k in sorted(hist, key=lambda k: float(k[1:].split(",")[0]))},
               "per_pdf": per_pdf}

    args.out.mkdir(parents=True, exist_ok=True)
    (args.out / "summary.json").write_text(json.dumps(summary, indent=1))
    with open(args.out / "structures.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(structs[0].keys()) if structs else ["file_name"])
        w.writeheader()
        for s in sorted(structs, key=lambda x: (x["pdf"], x["figure"] or 0, x["segment"] or 0)):
            w.writerow(s)
    md = [f"# Run summary: {args.label or args.run_dir.name}", "",
          f"{summary['pdfs']} PDFs, {summary['pages']} pages -> {summary['figures']} figures -> {summary['segments']} segments -> "
          f"{summary['structures']} structures ({summary['high']} high, {summary['low']} low; {summary['invalid']} unparseable).", "",
          "| PDF | pages | pages w/ figures | figures | segments | structures | high | low | invalid | median conf |", "|---|---|---|---|---|---|---|---|---|---|"]
    for pdf, v in per_pdf.items():
        md.append(f"| {pdf} | {v['pages']} | {v['pages_with_figures']} | {v['figures']} | {v['segments']} | {v['structures']} | {v['high']} | {v['low']} | {v['invalid']} | "
                  f"{None if v['median_confidence'] is None else round(v['median_confidence'], 3)} |")
    md += ["", "Confidence histogram (all structures):", ""]
    for k, v in summary["confidence_histogram"].items():
        md.append(f"- {k}: {v}")
    md += ["", "Structures per page (page numbers 1-based; pages absent produced no structure):", ""]
    for pdf, v in per_pdf.items():
        md.append(f"- {pdf}: " + ", ".join(f"p{k}:{n}" for k, n in v["structures_per_page"].items()))
    (args.out / "summary.md").write_text("\n".join(md) + "\n")
    print("\n".join(md))


if __name__ == "__main__":
    main()
