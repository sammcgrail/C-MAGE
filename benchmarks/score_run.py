#!/usr/bin/env python3
"""Score a C-MAGE run directory against a ground-truth manifest.

Reads the two stage-3 spreadsheets (Completed_HighConfidence_CMAGE.xlsx and
Completed_LowConfidence_CMAGE.xlsx), maps every emitted structure back to its
source group (PDF stem or image stem, carried in the segment file name), and
compares predicted and expected SMILES with RDKit canonical forms -- never by
string equality.

Verdict per emitted structure:
  exact    canonical isomeric SMILES equal to a molecule drawn in that group
  stereo   equal after dropping stereochemistry (right connectivity, stereo off)
  wrong    parses, but is not any molecule drawn in that group
  invalid  RDKit cannot parse it (includes the pipeline's literal "<invalid>")
  no-truth the group is not in the manifest (reported, excluded from precision)

Two denominators, reported separately and never merged:
  recall    = drawn molecules recovered / drawn molecules in the manifest
  precision = correct structures / structures the pipeline emitted for manifest groups
Both are given with stereochemistry required (exact) and relaxed (exact+stereo).

Correctness is cross-tabulated against the pipeline's own confidence split
(threshold 0.8431: >= is "high", else "low"; "<invalid>" is always low).

Outputs in --out:
  structures.csv   one row per emitted structure (verdict, confidence, tier, closest molecule, Tanimoto)
  molecules.csv    one row per drawn molecule (recovered exact / stereo / no)
  summary.json     all totals, per-group table, confidence cross-tab
  summary.md       the same as a readable table
  verdicts.json    per-group structures and molecules with image file names, for UIs
  sheets/<group>.png   with --sheets: segment image | render of prediction | render of expected, for hand review

Runs under the stage-3 environment (RDKit, pandas, openpyxl, Pillow are all there).

  score_run.py --run-dir results/run_X --manifest ground_truth/pdf_manifest.json --out results/scored_X --sheets
"""
import argparse
import csv
import datetime as dt
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from rdkit import Chem, DataStructs, RDLogger
from rdkit.Chem import AllChem, Draw

RDLogger.DisableLog("rdApp.*")

DIS_PREFIX = "Image_DIS_VH_File_"
THRESHOLD = 0.8431
HIGH = "high"
LOW = "low"


def parse_mol(smiles):
    """RDKit mol or None. Accepts CXSMILES; falls back to the plain SMILES part."""
    if smiles is None:
        return None
    s = str(smiles).strip()
    if not s or s == "<invalid>":
        return None
    m = Chem.MolFromSmiles(s)
    if m is None and "|" in s:
        m = Chem.MolFromSmiles(s.split("|", 1)[0].strip())
    return m


def canon(mol, stereo=True):
    return Chem.MolToSmiles(mol, isomericSmiles=stereo) if mol is not None else None


def group_key(path, groups=()):
    """Segment or image file -> (group, figure index, segment index).

    Stage 2 names segments Image_DIS_VH_File_<figure stem>_molecule_<k>.png and
    stage 1 names figures <pdf stem>_image_<n>.png. For an image corpus the figure
    stem is the image stem itself. A raw image path (stage-3-only run) has neither.

    Resolve against the manifest's own group names rather than the corpus LABEL.
    The `_image_<n>` strip used to be gated on `corpus == "pdfs"`, a literal string
    equality, and renaming the expanded PDF manifest's corpus to `pdfs_expanded`
    (needed so it stopped colliding with the original in the server's corpora dict)
    silently turned that gate off: every structure landed in a group named
    `EP0641330B1_pregabalin_image_4`, matched no manifest entry, was counted as
    "no ground truth" and excluded, and the run scored a clean 0/34 with no error
    and no warning -- a failure indistinguishable from an honest null result.
    Matching real group names cannot be broken by renaming a corpus.
    """
    stem = Path(path).stem
    if stem.startswith(DIS_PREFIX):
        stem = stem[len(DIS_PREFIX):]
    seg = fig = None
    m = re.match(r"^(.*)_molecule_(\d+)$", stem)
    if m:
        stem, seg = m.group(1), int(m.group(2))
    known = set(groups or ())
    if stem not in known:                       # an image corpus' stem IS the group
        m = re.match(r"^(.*)_image_(\d+)$", stem)
        if m and (not known or m.group(1) in known):
            stem, fig = m.group(1), int(m.group(2))
    return stem, fig, seg


def read_sheet(path, tier):
    """Rows of a stage-3 spreadsheet: (file path, predicted SMILES, confidence, classification)."""
    rows = []
    if not path.exists():
        return rows
    ws = load_workbook(path, read_only=True).active
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r is None or len(r) < 10 or r[1] is None:
            continue
        rows.append({"file": str(r[1]), "smiles": None if r[5] is None else str(r[5]),
                     "confidence": None if r[8] is None else float(r[8]),
                     "classification": r[9], "tier": tier})
    return rows


def render(mol, size):
    if mol is None:
        from PIL import Image, ImageDraw
        im = Image.new("RGB", size, "white")
        ImageDraw.Draw(im).text((8, 8), "(unparseable)", fill="red")
        return im
    return Draw.MolToImage(mol, size=size)


def make_sheet(group, structs, expected, out_path, run_dir):
    """One row per emitted structure: segment | prediction | expected (matched or closest)."""
    from PIL import Image, ImageDraw
    S = 300
    rows = []
    for s in structs:
        seg_path = Path(s["file"])
        if not seg_path.exists():
            seg_path = run_dir / "02_DIS_Segments" / seg_path.name
        try:
            seg = Image.open(seg_path).convert("RGB")
            seg.thumbnail((S, S))
        except Exception:  # noqa: BLE001
            seg = Image.new("RGB", (S, S), "grey")
        pred = render(parse_mol(s["smiles"]), (S, S))
        exp_name = s["matched_name"] or s["closest_name"]
        exp_mol = next((e["mol"] for e in expected if e["name"] == exp_name), None)
        exp = render(exp_mol, (S, S))
        row = Image.new("RGB", (3 * S + 20, S + 40), "white")
        row.paste(seg, (0, 40))
        row.paste(pred, (S + 10, 40))
        row.paste(exp, (2 * S + 20, 40))
        d = ImageDraw.Draw(row)
        d.text((2, 2), f"{Path(s['file']).name}  verdict={s['verdict']}  tier={s['tier']}  conf={s['confidence']}", fill="black")
        d.text((2, 20), f"pred: {str(s['smiles'])[:70]}", fill="blue")
        d.text((2 * S + 22, 20), f"expected: {exp_name}" + ("" if s["matched_name"] else f"  (closest, Tanimoto {s['closest_tanimoto']})"), fill="darkgreen")
        rows.append(row)
    if not rows:
        return
    sheet = Image.new("RGB", (rows[0].size[0], sum(r.size[1] for r in rows) + 20), "white")
    ImageDraw.Draw(sheet).text((2, 2), f"group {group}: expected " + ", ".join(e["name"] for e in expected), fill="black")
    y = 20
    for r in rows:
        sheet.paste(r, (0, y))
        y += r.size[1]
    sheet.save(out_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--run-dir", type=Path, required=True)
    ap.add_argument("--manifest", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--threshold", type=float, default=THRESHOLD)
    ap.add_argument("--sheets", action="store_true", help="write hand-review contact sheets")
    ap.add_argument("--label", default=None, help="free-text label for this run in summary")
    args = ap.parse_args()

    manifest = json.loads(args.manifest.read_text())
    corpus = manifest.get("corpus", "pdfs")
    expected = {}
    for g, entry in manifest["groups"].items():
        mols = []
        for m in entry["molecules"]:
            mol = Chem.MolFromSmiles(m["smiles"])
            if mol is None:
                raise SystemExit(f"manifest SMILES for {g}/{m.get('name')} does not parse")
            mols.append({"name": m.get("name") or m.get("label"), "label": m.get("label", m.get("name")),
                         "smiles": m["smiles"], "mol": mol, "iso": canon(mol, True), "flat": canon(mol, False),
                         "heavy": mol.GetNumHeavyAtoms(),
                         "fp": AllChem.GetMorganFingerprintAsBitVect(mol, 2, 2048)})
        expected[g] = {"source": entry.get("source"), "molecules": mols}

    res = args.run_dir / "03_CXMS_Results"
    rows = read_sheet(res / "Completed_HighConfidence_CMAGE.xlsx", HIGH) + \
        read_sheet(res / "Completed_LowConfidence_CMAGE.xlsx", LOW)
    if not rows:
        raise SystemExit(f"no stage-3 rows found under {res}")

    # segments stage 2 handed to stage 3 (to catch any that produced no row)
    seg_list = []
    seg_xlsx = args.run_dir / "02_DIS_Segments" / "DIS_CMAGE_results.xlsx"
    if seg_xlsx.exists():
        ws = load_workbook(seg_xlsx, read_only=True).active
        header = None
        for r in ws.iter_rows(values_only=True):
            if header is None:
                header = list(r)
                continue
            if r and r[-1]:
                seg_list.append(str(r[-1]))
    figures = sorted(p.name for p in (args.run_dir / "01_VH_Figures").glob("*.png")) if (args.run_dir / "01_VH_Figures").exists() else []

    structs = []
    for r in rows:
        g, fig, seg = group_key(r["file"], expected)
        mol = parse_mol(r["smiles"])
        rec = {"group": g, "figure": fig, "segment": seg, "file": r["file"], "file_name": Path(r["file"]).name,
               "smiles": r["smiles"], "confidence": r["confidence"], "tier": r["tier"],
               "tier_by_threshold": (HIGH if (r["confidence"] is not None and r["confidence"] >= args.threshold and mol is not None) else LOW),
               "classification": r["classification"], "parsed": mol is not None,
               "pred_heavy": mol.GetNumHeavyAtoms() if mol else None,
               "verdict": None, "matched_name": None, "closest_name": None, "closest_tanimoto": None, "closest_heavy": None}
        if g not in expected:
            rec["verdict"] = "no-truth"
        elif mol is None:
            rec["verdict"] = "invalid"
        else:
            iso, flat = canon(mol, True), canon(mol, False)
            exp = expected[g]["molecules"]
            hit = next((e for e in exp if e["iso"] == iso), None)
            if hit:
                rec["verdict"], rec["matched_name"] = "exact", hit["name"]
            else:
                hit = next((e for e in exp if e["flat"] == flat), None)
                if hit:
                    rec["verdict"], rec["matched_name"] = "stereo", hit["name"]
                else:
                    rec["verdict"] = "wrong"
            fp = AllChem.GetMorganFingerprintAsBitVect(mol, 2, 2048)
            sims = [(DataStructs.TanimotoSimilarity(fp, e["fp"]), e) for e in exp]
            best = max(sims, key=lambda t: t[0])
            rec["closest_name"], rec["closest_tanimoto"], rec["closest_heavy"] = best[1]["name"], round(best[0], 3), best[1]["heavy"]
        structs.append(rec)

    # ---- recall: per drawn molecule
    mol_rows = []
    for g, entry in expected.items():
        in_group = [s for s in structs if s["group"] == g]
        for e in entry["molecules"]:
            ex = [s for s in in_group if s["verdict"] == "exact" and s["matched_name"] == e["name"]]
            st = [s for s in in_group if s["verdict"] == "stereo" and s["matched_name"] == e["name"]]
            rec = "exact" if ex else ("stereo" if st else "no")
            mol_rows.append({"group": g, "source": entry["source"], "name": e["name"], "label": e["label"],
                             "heavy_atoms": e["heavy"], "recovered": rec,
                             "structures_in_group": len(in_group),
                             "best_confidence": max([s["confidence"] or 0 for s in ex + st], default=None),
                             "matching_structures": len(ex) + len(st)})

    scored = [s for s in structs if s["verdict"] != "no-truth"]
    no_truth = [s for s in structs if s["verdict"] == "no-truth"]

    def count(lst, **kw):
        return sum(1 for s in lst if all(s[k] == v for k, v in kw.items()))

    n_mols = len(mol_rows)
    n_struct = len(scored)
    summary = {
        "label": args.label, "run_dir_name": args.run_dir.name, "manifest": args.manifest.name, "corpus": corpus,
        "generated": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"), "threshold": args.threshold,
        "figures_from_stage1": len(figures), "segments_from_stage2": len(seg_list), "structures_emitted_total": len(structs),
        "structures_for_groups_without_truth": len(no_truth),
        "drawn_molecules": n_mols,
        "recall": {
            "denominator": n_mols,
            "recovered_exact": sum(1 for m in mol_rows if m["recovered"] == "exact"),
            "recovered_exact_or_stereo": sum(1 for m in mol_rows if m["recovered"] in ("exact", "stereo")),
            "not_recovered": sum(1 for m in mol_rows if m["recovered"] == "no"),
            "groups_with_no_structure_at_all": sum(1 for g in expected if not any(s["group"] == g for s in structs)),
        },
        "precision": {
            "denominator": n_struct,
            "exact": count(scored, verdict="exact"), "stereo": count(scored, verdict="stereo"),
            "wrong": count(scored, verdict="wrong"), "invalid": count(scored, verdict="invalid"),
        },
        "confidence_crosstab": {},
        "tier_disagreements_with_threshold": sum(1 for s in scored if s["tier"] != s["tier_by_threshold"]),
        "per_group": [],
    }
    r, p = summary["recall"], summary["precision"]
    r["recall_exact"] = round(r["recovered_exact"] / n_mols, 4) if n_mols else None
    r["recall_exact_or_stereo"] = round(r["recovered_exact_or_stereo"] / n_mols, 4) if n_mols else None
    p["precision_exact"] = round(p["exact"] / n_struct, 4) if n_struct else None
    p["precision_exact_or_stereo"] = round((p["exact"] + p["stereo"]) / n_struct, 4) if n_struct else None
    for tier in (HIGH, LOW):
        t = [s for s in scored if s["tier"] == tier]
        c = {v: count(t, verdict=v) for v in ("exact", "stereo", "wrong", "invalid")}
        c["total"] = len(t)
        c["fraction_wrong_or_invalid"] = round((c["wrong"] + c["invalid"]) / len(t), 4) if t else None
        c["fraction_exact"] = round(c["exact"] / len(t), 4) if t else None
        c["fraction_exact_or_stereo"] = round((c["exact"] + c["stereo"]) / len(t), 4) if t else None
        summary["confidence_crosstab"][tier] = c
    for g, entry in expected.items():
        in_group = [s for s in structs if s["group"] == g]
        mols = [m for m in mol_rows if m["group"] == g]
        summary["per_group"].append({
            "group": g, "source": entry["source"], "drawn": len(mols),
            "figures": sum(1 for f in figures if group_key(f, expected)[0] == g) if figures else None,
            "segments": sum(1 for f in seg_list if group_key(f, expected)[0] == g) if seg_list else None,
            "structures": len(in_group),
            "recovered_exact": sum(1 for m in mols if m["recovered"] == "exact"),
            "recovered_stereo": sum(1 for m in mols if m["recovered"] == "stereo"),
            "exact": count(in_group, verdict="exact"), "stereo": count(in_group, verdict="stereo"),
            "wrong": count(in_group, verdict="wrong"), "invalid": count(in_group, verdict="invalid"),
            "high": count(in_group, tier=HIGH), "low": count(in_group, tier=LOW),
        })

    # ---- write outputs
    args.out.mkdir(parents=True, exist_ok=True)
    s_fields = ["group", "figure", "segment", "file_name", "tier", "tier_by_threshold", "confidence", "verdict", "matched_name",
                "closest_name", "closest_tanimoto", "pred_heavy", "closest_heavy", "smiles"]
    with open(args.out / "structures.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=s_fields, extrasaction="ignore")
        w.writeheader()
        for s in sorted(structs, key=lambda x: (x["group"], x["figure"] or 0, x["segment"] or 0)):
            w.writerow(s)
    with open(args.out / "molecules.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(mol_rows[0].keys()) if mol_rows else ["group"])
        w.writeheader()
        for m in mol_rows:
            w.writerow(m)
    (args.out / "summary.json").write_text(json.dumps(summary, indent=1))

    # verdicts.json for UIs: image file names only, no absolute paths
    list_key = "pdfs" if corpus == "pdfs" else "images"
    verdicts = {"corpus": corpus, "threshold": args.threshold, "verdict_values": ["exact", "stereo", "wrong", "invalid"],
                list_key: []}
    for g, entry in expected.items():
        verdicts[list_key].append({
            "group": g, "file": entry["source"],
            "molecules": [{"key": e["name"], "label": e["label"], "smiles": e["smiles"]} for e in entry["molecules"]],
            "structures": [{"segment_image": s["file_name"],
                            "rendered_image": (Path(s["file_name"]).stem[len(DIS_PREFIX):] if s["file_name"].startswith(DIS_PREFIX) else Path(s["file_name"]).stem) + ".png",
                            "rendered_dir": "highconfidence_images" if s["tier"] == HIGH else "lowconfidence_images",
                            "smiles": s["smiles"], "confidence": s["confidence"], "tier": s["tier"],
                            "verdict": s["verdict"], "matched_key": s["matched_name"],
                            "closest_key": s["closest_name"], "closest_tanimoto": s["closest_tanimoto"]}
                           for s in structs if s["group"] == g]})
    if no_truth:
        verdicts["structures_without_truth"] = [{"group": s["group"], "segment_image": s["file_name"], "smiles": s["smiles"],
                                                 "confidence": s["confidence"], "tier": s["tier"]} for s in no_truth]
    (args.out / "verdicts.json").write_text(json.dumps(verdicts, indent=1))

    md = [f"# Score: {args.label or args.run_dir.name}", "",
          f"Manifest `{args.manifest.name}` ({corpus}); threshold {args.threshold}.", "",
          f"- Drawn molecules (recall denominator): **{n_mols}**",
          f"- Structures emitted for those groups (precision denominator): **{n_struct}**"
          + (f" (+{len(no_truth)} for groups without ground truth, excluded)" if no_truth else ""),
          f"- Stage 1 figures: {len(figures) or 'n/a'}; stage 2 segments: {len(seg_list) or 'n/a'}", "",
          "| metric | stereo required | stereo relaxed |", "|---|---|---|",
          f"| recall | {r['recovered_exact']}/{n_mols} = {r['recall_exact']} | {r['recovered_exact_or_stereo']}/{n_mols} = {r['recall_exact_or_stereo']} |",
          f"| precision | {p['exact']}/{n_struct} = {p['precision_exact']} | {p['exact'] + p['stereo']}/{n_struct} = {p['precision_exact_or_stereo']} |",
          "", "Structures by verdict: " + ", ".join(f"{k} {p[k]}" for k in ("exact", "stereo", "wrong", "invalid")), "",
          "## Confidence split vs correctness", "",
          "| tier | n | exact | stereo | wrong | invalid | wrong+invalid | exact |", "|---|---|---|---|---|---|---|---|"]
    for tier in (HIGH, LOW):
        c = summary["confidence_crosstab"][tier]
        md.append(f"| {tier} | {c['total']} | {c['exact']} | {c['stereo']} | {c['wrong']} | {c['invalid']} | {c['fraction_wrong_or_invalid']} | {c['fraction_exact']} |")
    md += ["", "## Per group", "", "| group | drawn | figs | segs | structs | rec.exact | rec.stereo | exact | stereo | wrong | invalid | high | low |",
           "|---|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for g in summary["per_group"]:
        md.append("| " + " | ".join(str(g[k]) for k in ("group", "drawn", "figures", "segments", "structures", "recovered_exact",
                                                          "recovered_stereo", "exact", "stereo", "wrong", "invalid", "high", "low")) + " |")
    (args.out / "summary.md").write_text("\n".join(md) + "\n")

    if args.sheets:
        sd = args.out / "sheets"
        sd.mkdir(exist_ok=True)
        for g, entry in expected.items():
            in_group = sorted([s for s in structs if s["group"] == g], key=lambda x: (x["figure"] or 0, x["segment"] or 0))
            if in_group:
                make_sheet(g, in_group, entry["molecules"], sd / f"{g}.png", args.run_dir)

    print("\n".join(md))


if __name__ == "__main__":
    main()
