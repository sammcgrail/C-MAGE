#!/usr/bin/env python3
"""Extract stage timings and throughput from a C-MAGE run directory.

Sources, in order of trust:
  logs/stage1.log   "Total VisualHeist Time = N seconds"            (includes model load)
  logs/stage2.log   "Total DECIMER-Image-Segmentation Time = N"     (excludes model load; timer starts after import)
  logs/stage3.log   "Total CXMolScribe Time = N seconds"            (excludes model load and spreadsheet writing)
  --time-log FILE   output of `/usr/bin/time -v` wrapped around the whole run: wall clock and max RSS
  file counts       01_VH_Figures/*.png, 02_DIS_Segments/*.png, rows in the two stage-3 spreadsheets
  --pdf-dir DIR     page count of the inputs via pdfinfo, for per-page rates

Prints a JSON object and a one-line markdown table row.

  timings.py --run-dir results/run_X [--time-log run.log] [--pdf-dir corpus/] [--label "ground-truth PDFs"]
"""
import argparse
import json
import re
import subprocess
from pathlib import Path


def grab(path, pattern):
    if not path.exists():
        return None
    m = re.search(pattern, path.read_text(errors="replace"))
    return float(m.group(1)) if m else None


def time_v(path):
    if not path or not path.exists():
        return {}
    t = path.read_text(errors="replace")
    out = {}
    m = re.search(r"Elapsed \(wall clock\) time \(h:mm:ss or m:ss\):\s*([\d:.]+)", t)
    if m:
        parts = [float(x) for x in m.group(1).split(":")]
        out["wall_clock_s"] = round(sum(p * 60 ** i for i, p in enumerate(reversed(parts))), 1)
    m = re.search(r"Maximum resident set size \(kbytes\):\s*(\d+)", t)
    if m:
        out["max_rss_mb"] = round(int(m.group(1)) / 1024)
    m = re.search(r"Percent of CPU this job got:\s*(\d+)%", t)
    if m:
        out["cpu_percent"] = int(m.group(1))
    return out


def pages(pdf_dir):
    if not pdf_dir or not pdf_dir.exists():
        return None
    n = 0
    for pdf in pdf_dir.glob("*.pdf"):
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
        m = re.search(r"^Pages:\s+(\d+)", info, re.M)
        n += int(m.group(1)) if m else 0
    return n


def rows_in(xlsx):
    if not xlsx.exists():
        return 0
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, read_only=True).active
    return sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r and r[1] is not None)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--run-dir", type=Path, required=True)
    ap.add_argument("--time-log", type=Path, default=None)
    ap.add_argument("--pdf-dir", type=Path, default=None)
    ap.add_argument("--label", default="")
    args = ap.parse_args()
    rd = args.run_dir
    logs = rd / "logs"
    figs = len(list((rd / "01_VH_Figures").glob("*.png"))) if (rd / "01_VH_Figures").exists() else None
    segs = len(list((rd / "02_DIS_Segments").glob("*.png"))) if (rd / "02_DIS_Segments").exists() else None
    res = rd / "03_CXMS_Results"
    structs = rows_in(res / "Completed_HighConfidence_CMAGE.xlsx") + rows_in(res / "Completed_LowConfidence_CMAGE.xlsx")
    out = {
        "label": args.label, "run": rd.name,
        "pages": pages(args.pdf_dir),
        "stage1_s": grab(logs / "stage1.log", r"Total VisualHeist Time = ([\d.]+)"),
        "stage2_s": grab(logs / "stage2.log", r"Total DECIMER-Image-Segmentation Time = ([\d.]+)"),
        "stage3_s": grab(logs / "stage3.log", r"Total CXMolScribe Time = ([\d.]+)"),
        "figures": figs, "segments": segs, "structures": structs,
        "stage2_no_segment_figures": len(re.findall(r"FAILED NO SEGMENT", (logs / "stage2.log").read_text(errors="replace"))) if (logs / "stage2.log").exists() else None,
    }
    out.update(time_v(args.time_log))
    if out["stage1_s"] and out["pages"]:
        out["stage1_s_per_page"] = round(out["stage1_s"] / out["pages"], 1)
    if out["stage2_s"] and figs:
        out["stage2_s_per_figure"] = round(out["stage2_s"] / figs, 1)
    if out["stage3_s"] and structs:
        out["stage3_s_per_structure"] = round(out["stage3_s"] / structs, 2)
    print(json.dumps(out, indent=1))
    print("| " + " | ".join(str(out.get(k)) for k in ("label", "pages", "figures", "segments", "structures", "stage1_s", "stage2_s", "stage3_s", "wall_clock_s", "max_rss_mb")) + " |")


if __name__ == "__main__":
    main()
